import html
import json
import os
import shutil
import tempfile
from pathlib import Path
from PIL import Image
from playwright.sync_api import sync_playwright
from pptx import Presentation
from pptx.util import Inches
from .schemas import validate
from .theme import load_theme
from .renderer import render_html, SOURCE, ROOT

def launch_browser(playwright):
    executable = os.environ.get('PPTXGEN_CHROMIUM')
    return playwright.chromium.launch(executable_path=executable) if executable else playwright.chromium.launch()

def inspect_page(page):
    page.evaluate('document.fonts.ready')
    page.evaluate('async () => { await Promise.all([...document.images].map(i => i.decode().catch(() => {}))); }')
    return page.evaluate((SOURCE / 'qa.js').read_text())

def _copy_review_artifacts(work, output):
    for directory in ('html', 'images'):
        target = output / directory; target.mkdir(exist_ok=True)
        for old in target.glob('slide_*.' + ('html' if directory == 'html' else 'png')): old.unlink()
        shutil.copytree(work / directory, target, dirs_exist_ok=True)
    shutil.copy2(work / 'preview.html', output / 'preview.html')

def _make_pptx(images, target, dimensions):
    width, height = dimensions
    deck = Presentation(); deck.slide_width, deck.slide_height = Inches(width / 144), Inches(height / 144)
    for png in images:
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        slide.shapes.add_picture(str(png), 0, 0, width=deck.slide_width, height=deck.slide_height)
    deck.save(target)
    if len(Presentation(target).slides) != len(images): raise ValueError('PPTX_COUNT: slide count mismatch')

def generate(input_path, output, theme_name='default', ratio='16:9', root=ROOT, approve=False):
    """Render and QA HTML/PNG. PPTX is created only when approve=True."""
    output = Path(output).resolve(); output.mkdir(parents=True, exist_ok=True)
    report = {'status': 'error', 'theme': theme_name, 'slides': [], 'issues': [], 'generated': False, 'awaiting_human_review': False}
    try:
        data = json.loads(Path(input_path).read_text()); validate(data)
        logo, theme = load_theme(Path(root), theme_name); report['logo_source'] = str(logo.path)
        width, height = (1920, 1080) if ratio == '16:9' else (1440, 1080)
        with tempfile.TemporaryDirectory(prefix='pptxgen-') as temporary:
            work = Path(temporary); (work/'html').mkdir(); (work/'images').mkdir()
            with sync_playwright() as pw:
                browser = launch_browser(pw)
                try:
                    page = browser.new_page(viewport={'width': width, 'height': height}, device_scale_factor=1)
                    for slide in data['slides']:
                        stem = f"slide_{slide['slide_number']:03}"; path = work/'html'/f'{stem}.html'
                        path.write_text(render_html(slide, data['presentation'], len(data['slides']), logo, theme, width, height))
                        page.goto(path.as_uri()); result = inspect_page(page)
                        report['slides'].append({'slide_number': slide['slide_number'], **result})
                        if result['issues']: continue
                        png = work/'images'/f'{stem}.png'; page.screenshot(path=str(png), full_page=False)
                        with Image.open(png) as img:
                            if img.size != (width, height): raise ValueError(f'PNG_SIZE: {img.size}')
                finally: browser.close()
            if any(s['issues'] for s in report['slides']): raise ValueError('SLIDE_QA_FAILED: see slides in qa_report.json')
            images = sorted((work/'images').glob('*.png'))
            if len(images) != len(data['slides']): raise ValueError('SLIDE_COUNT: missing rendered slides')
            links = ''.join(f'<a href="images/{p.name}"><img src="images/{p.name}" alt="Slide {i}"></a>' for i, p in enumerate(images, 1))
            (work/'preview.html').write_text('<!doctype html><html lang="ja"><meta charset="utf-8"><title>Presentation Preview</title>'
                '<style>body{background:#e9eef2;font-family:sans-serif;margin:40px}main{display:grid;grid-template-columns:repeat(auto-fit,minmax(320px,1fr));gap:24px}img{width:100%}</style>'
                f'<h1>{html.escape(data["presentation"]["title"])}</h1><main>{links}</main></html>')
            _copy_review_artifacts(work, output)
            (output/'review_manifest.json').write_text(json.dumps({'input': str(Path(input_path).resolve()), 'theme': theme_name, 'ratio': ratio, 'dimensions': [width, height], 'slide_count': len(images)}, ensure_ascii=False, indent=2))
            if approve:
                _make_pptx(images, work/'presentation.pptx', (width, height)); shutil.copy2(work/'presentation.pptx', output/'presentation.pptx')
                report.update(status='passed', generated=True, slide_count=len(images), dimensions=[width, height], approved=True)
            else:
                report.update(status='awaiting_human_review', awaiting_human_review=True, slide_count=len(images), dimensions=[width, height])
    except Exception as exc: report['issues'].append({'severity': 'error', 'code': 'GENERATION_ERROR', 'detail': str(exc)})
    finally: (output/'qa_report.json').write_text(json.dumps(report, ensure_ascii=False, indent=2))
    return report

def approve_html(output):
    """Package the already reviewed HTML render; never regenerates HTML."""
    output = Path(output).resolve(); manifest = json.loads((output/'review_manifest.json').read_text()); report = json.loads((output/'qa_report.json').read_text())
    if report.get('status') != 'awaiting_human_review' or report.get('issues'): raise ValueError('Approval requires a passed HTML review report')
    images = sorted((output/'images').glob('slide_*.png'))
    if len(images) != manifest['slide_count']: raise ValueError('APPROVAL_SLIDE_COUNT')
    _make_pptx(images, output/'presentation.pptx', manifest['dimensions'])
    report.update(status='passed', generated=True, approved=True); (output/'qa_report.json').write_text(json.dumps(report, ensure_ascii=False, indent=2))
    return report
