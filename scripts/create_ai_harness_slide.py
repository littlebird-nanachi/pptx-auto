from pathlib import Path
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'output' / 'ai_harness_4x3.pptx'
LOGO = ROOT / 'assets' / 'logo.png'

NAVY = RGBColor(31, 48, 59); RED = RGBColor(169, 29, 42); PALE_RED = RGBColor(248, 232, 233)
GRAY = RGBColor(235, 239, 241); MID = RGBColor(88, 101, 109); WHITE = RGBColor(255,255,255)
BLACK = RGBColor(24, 35, 42); LIGHT = RGBColor(248, 249, 250)

prs = Presentation(); prs.slide_width = Inches(10); prs.slide_height = Inches(7.5)
slide = prs.slides.add_slide(prs.slide_layouts[6]); slide.background.fill.solid(); slide.background.fill.fore_color.rgb = WHITE

def box(x,y,w,h,fill=None,line=None,radius=False):
    shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shp.fill.solid(); shp.fill.fore_color.rgb = fill or WHITE
    shp.line.color.rgb = line or (fill or WHITE); shp.line.width = Pt(0.7)
    return shp

def text(x,y,w,h,s,size=9,color=BLACK,bold=False,align=PP_ALIGN.LEFT,margin=0.05):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)); tf=tb.text_frame; tf.clear(); tf.word_wrap=True
    tf.margin_left=tf.margin_right=Inches(margin); tf.margin_top=tf.margin_bottom=Inches(margin); tf.vertical_anchor=MSO_ANCHOR.MIDDLE
    p=tf.paragraphs[0]; p.alignment=align; p.space_after=Pt(0); p.space_before=Pt(0)
    r=p.add_run(); r.text=s; r.font.name='Arial'; r.font.size=Pt(size); r.font.bold=bold; r.font.color.rgb=color
    return tb

def arrow(x1,y1,x2,y2,color=MID,width=1.1):
    ln=slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2)); ln.line.color.rgb=color; ln.line.width=Pt(width); ln.line.end_arrowhead=True; return ln

# Header
text(.45,.28,8.2,.38,'AI活用ポイントとInput / Output',24,NAVY,True)
text(.45,.73,8.55,.28,'生成AIを5つの機能ブロックに分け、Rule-based処理とHuman Gateとの責任分界を明確化',11,MID)
if LOGO.exists(): slide.shapes.add_picture(str(LOGO), Inches(8.85), Inches(.25), height=Inches(.34))

cards = [
 ('AI-01','顧客要求の構造化','RFI / 物理アーキ\nExcel・PDF','要求・制約を抽出\n認識・JSON構造化','Requirement JSON','Rule：Validation / 単位統一','Human：要求FIX'),
 ('AI-02','Type X差分解釈','要求 + 標準仕様\n+ Gap','差分解釈・対応\n候補提示','Option / 新規候補\n＋根拠','Rule：比較 / Gap算出','Human：方針FIX'),
 ('AI-03','電気設計候補探索','Gap + 部品 /\n標準設計資産','解消候補を探索\n・整理','電気設計候補','Rule：Heat / PMIC / Area','Human：電気案採用'),
 ('AI-04','機械候補・Feedback','電気設計結果 +\n機械制約','配置・改善候補\nを提示','機械候補 /\n電気Feedback','Rule：寸法 / 設備判定','Human：機械案採用'),
 ('AI-05','設計結果の要約','全工程の確定結果\n/ Decision Log','判断・根拠・履歴\nを要約','Design Study\nReport','Rule：State / STALE確認','Human：Final Approval'),
]

def card(x,y,w,h,item):
    no,title,inp,proc,out,rule,human=item
    box(x,y,w,h,LIGHT,RGBColor(190,198,202),True)
    box(x,y,w,.34,RED,RED,True); text(x+.08,y+.01,w-.16,.28,no+'  '+title,10,WHITE,True)
    yy=y+.43
    for label,val,col in [('INPUT',inp,NAVY),('AI PROCESS',proc,RED),('OUTPUT',out,NAVY)]:
        text(x+.1,yy,w-.2,.16,label,8,col,True); text(x+.1,yy+.16,w-.2,.43,val,9,BLACK,False); yy += .66
    box(x+.08,y+h-.83,w-.16,.3,GRAY,GRAY,True); text(x+.15,y+h-.81,w-.3,.25,rule,8.2,MID,True)
    box(x+.08,y+h-.47,w-.16,.3,PALE_RED,RED,True); text(x+.15,y+h-.45,w-.3,.25,human,8.2,RED,True)

# Two-row cards keep 4:3 readable.
xs=[.45,2.37,4.29,6.21,8.13]; w=1.7
for i,item in enumerate(cards): card(xs[i],1.28 if i<3 else 3.98,w,2.32,item)
arrow(2.16,2.43,2.34,2.43); arrow(4.08,2.43,4.26,2.43)
arrow(7.92,5.13,8.1,5.13)
text(.48,3.67,5.7,.2,'要求入力  →  Type X判定  →  電気設計',8.5,MID,True)
text(6.23,6.37,3.0,.2,'機械設計  →  最終設計・要約',8.5,MID,True)
# Feedback loop marker for AI-04
arrow(7.05,6.31,5.2,6.31,RED,1.0); text(5.27,6.12,1.7,.18,'Mechanical NG → Electrical Feedback',7.5,RED,True)

# Bottom responsibility strip
roles=[('AI','読む / 探す / 解釈する / 候補を出す / 要約する',PALE_RED,RED),('Rule-based','計算 / 比較 / 制約判定 / 状態管理 / Validation',GRAY,MID),('Human','確認 / 修正 / 選択 / 設計判断 / 承認',PALE_RED,RED)]
for i,(name,desc,fill,line) in enumerate(roles):
    x=.45+i*3.04; box(x,6.68,2.88,.42,fill,line,True); text(x+.1,6.71,.72,.18,name,9.5,line,True); text(x+.8,6.71,1.95,.18,desc,7.3,BLACK)
text(.45,7.15,9.1,.18,'AIは候補と判断材料を生成し、Rule-basedが成立性を保証し、Humanが最終判断する',9.5,NAVY,True)

OUT.parent.mkdir(exist_ok=True); prs.save(OUT); print(OUT)
